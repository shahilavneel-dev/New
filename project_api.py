"""Local project/settings API and confidential Excel/PDF estimating reports."""
from __future__ import annotations

import copy
import io
import math
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import fitz
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field

from storage import RevisionConflict, Store, import_sections, json_text, normal_section, validate_settings


class StateRequest(BaseModel):
    state: dict[str, Any]


class SaveRequest(StateRequest):
    id: str | None = None
    expected_revision: int | None = Field(default=None, ge=0, strict=True)


def _calculate(state: dict, store: Store) -> dict:
    from estimating import calculate
    encoded = json_text(state)
    if len(encoded.encode("utf-8")) > 20 * 1024 * 1024:
        raise ValueError("Project JSON exceeds 20 MB")
    if not isinstance(state.get("project"), dict):
        raise ValueError("Project state must contain a project object")
    if not isinstance(state.get("members", []), list) or not isinstance(state.get("plates", []), list):
        raise ValueError("Members and plates must be arrays")
    if any(not isinstance(r, dict) for r in state.get("members", []) + state.get("plates", [])):
        raise ValueError("Each member and plate must be an object")
    settings = store.settings()
    overrides = state.get("rates", {})
    if not isinstance(overrides, dict):
        raise ValueError("Project rate overrides must be an object")
    effective = validate_settings({"rates": overrides}, settings)["rates"]
    prepared = copy.deepcopy(state)
    catalog = {normal_section(s["section"]): s for s in settings["sections"]}
    for member in prepared.get("members", []):
        entry = catalog.get(normal_section(member.get("section", ""))) if isinstance(member, dict) else None
        if not entry:
            continue
        # Manual project overrides take precedence over the shared catalogue.
        if member.get("kgm") in (None, ""):
            member["kgm"] = entry["kgm"]
        if not member.get("stocks") and entry.get("stocks"):
            member["stocks"] = ",".join(str(x) for x in entry["stocks"])
        if member.get("supplierRate") in (None, "") and entry.get("supplier_rate") is not None:
            member["supplierRate"] = entry["supplier_rate"]
    result = calculate(prepared, rates=effective)
    # Saved reports use this exact calculation snapshot rather than today's rates.
    result["effectiveRates"] = effective
    result["calculatedAt"] = datetime.now(timezone.utc).isoformat()
    result["resolvedMembers"] = prepared.get("members", [])
    return result


def _safe_name(name: str, suffix: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", name or "SteelTakeoff")[:65].strip("_") + suffix


def _numeric(value: Any) -> float:
    try:
        n = float(value)
        return n if math.isfinite(n) else 0
    except (ValueError, TypeError):
        return 0


def _money(value: Any) -> str:
    return "Unverified" if value is None else f"${_numeric(value):,.2f}"


def _text(value: Any) -> str:
    return "" if value is None else str(value)


def _sheet_value(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def _report_parts(record: dict):
    state = record["state"]
    result = record["result"]
    project = state.get("project", {})
    return state, result, project


def _drawing_review_notes(state: dict) -> list[list[str]]:
    """Carry drawing evidence into internal reports without implying verification."""
    rows = [["Drawing review flag", _text(flag)] for flag in state.get("pdfFlags", [])]
    cross = state.get("pdfCrossReferences")
    if not cross:
        return rows
    if not isinstance(cross, dict):
        rows.append(["Drawing cross-reference - review required", _text(cross)])
        return rows
    rows.append(["Drawing cross-reference review", _text(cross.get("review") or "Advisory evidence only; verify against the drawing.")])
    for mark in cross.get("marks", []):
        if not isinstance(mark, dict):
            rows.append(["Member cross-reference - review required", _text(mark)])
            continue
        label = _text(mark.get("mark"))
        conflict = "SECTION CONFLICT - review required" if mark.get("section_conflict") else "Cross-reference evidence - verify scope and quantity"
        rows.append([f"Member {label}: {conflict}", json_text(mark)])
    for suggestion in cross.get("rl_suggestions", []):
        rows.append(["PROVISIONAL RL height suggestion - verify end conditions", json_text(suggestion)])
    for flag in cross.get("flags", []):
        if _text(flag) not in {_text(row[1]) for row in rows}:
            rows.append(["Cross-reference review flag", _text(flag)])
    if cross.get("method"):
        rows.append(["Drawing cross-reference method", _text(cross["method"])])
    return rows


def excel_report(record: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    state, result, project = _report_parts(record)
    drawing_notes = _drawing_review_notes(state)
    density = result.get("effectiveRates", {}).get("steel_density_kg_mm3", result.get("rates", {}).get("steel_density_kg_mm3", .00000785))
    book = Workbook()
    book.remove(book.active)
    navy, teal = "172B47", "007F86"

    def sheet(name: str, headers: list, rows: list, widths=None):
        ws = book.create_sheet(name)
        ws.append([f"SteelTakeoff AI | INTERNAL | {project.get('name') or 'Untitled project'}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, len(headers)))
        ws.cell(1, 1).font = Font(size=15, color="FFFFFF", bold=True)
        ws.cell(1, 1).fill = PatternFill("solid", fgColor=navy)
        ws.append([f"PRIVATE - NOT FOR CUSTOMER ISSUE | Revision {record['revision']} | Saved {record.get('saved_at', '')} | AUD excluding GST"])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(2, len(headers)))
        ws.cell(2, 1).font = Font(size=9, color="68768A")
        ws.append(headers)
        for cell in ws[3]:
            cell.fill = PatternFill("solid", fgColor=teal)
            cell.font = Font(bold=True, color="FFFFFF")
        for row in rows:
            ws.append([_sheet_value(v) for v in row])
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="EFF4F8")
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"
        for i in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1] if widths and i <= len(widths) else 21
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{max(3,ws.max_row)}"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.print_title_rows = "1:3"
        ws.oddFooter.center.text = "INTERNAL - Page &P of &N"
        return ws

    summary = [
        ["Project", project.get("name", "")], ["Customer", project.get("client", "")],
        ["Address", project.get("address", "")], ["Scope", project.get("scope", "")],
        ["Net steel incl. plates (kg)", result.get("net", 0)],
        ["Gross purchased sections (kg)", result.get("purchase", {}).get("gross", 0)],
        ["Fabricated plate allowance (kg)", result.get("plateWeight", 0)],
        ["Charged plate material (kg)", result.get("plateGross", 0)],
        ["Internal cost (known costs only if incomplete)", result.get("totalCost", 0)],
        ["Supply & fabrication selling", result.get("supplySell", 0)],
        ["Installation selling", result.get("instSell", 0)],
        ["Total selling (ex GST)", result.get("totalSell", 0)],
        ["Gross profit", result.get("profit") if result.get("profit") is not None else "Unverified - incomplete costs"],
        ["Overall gross margin (%)", result.get("margin") if result.get("margin") is not None else "Unverified - incomplete costs"],
        ["Estimate status", "REVIEW WARNINGS BEFORE ISSUE" if result.get("warnings") or drawing_notes else "Review assumptions before issue"],
        ["Scope / assumptions", project.get("notes", "")],
    ]
    sheet("Summary", ["Metric", "Value"], summary, [48, 88])
    members = result.get("resolvedMembers", state.get("members", []))
    sheet("Members", ["Mark", "Family", "Section", "Quantity", "Length mm", "kg/m", "Net kg", "Finish", "Process", "Drawing", "Page", "Evidence / status"], [
        [m.get("mark", ""), m.get("family", ""), m.get("section", ""), m.get("qty", ""), m.get("length", ""), m.get("kgm", ""), _numeric(m.get("qty"))*_numeric(m.get("length"))*_numeric(m.get("kgm"))/1000, m.get("finish", ""), m.get("process", ""), m.get("sourceFile", ""), m.get("sourcePage", ""), m.get("sourceEvidence", m.get("status", ""))] for m in members
    ], [16, 13, 24, 12, 15, 12, 15, 14, 14, 32, 10, 60])
    allocations = result.get("purchase", {}).get("allocations", [])
    stock_rows = []
    for a in allocations:
        row = a.get("row", {})
        cuts = a.get("selected", [])
        stock_rows.append([row.get("section", ""), a.get("kind", "stock"), a.get("L", 0), " + ".join(f"{p.get('mark','')}:{p.get('length',0)}" for p in cuts), _numeric(a.get("L"))*_numeric(row.get("kgm")), a.get("cost", 0), a.get("remainingNominal", ""), "PROVISIONAL - confirm lengths/kerf" if a.get("provisional") else row.get("process", "")])
    sheet("Stock purchasing", ["Section", "Purchase", "Stock m", "Cut lengths mm", "Gross kg", "Internal buy cost", "Remainder mm", "Review"], stock_rows, [24, 18, 13, 65, 15, 20, 20, 45])
    layouts = book.create_sheet("Stock layouts")
    layouts.merge_cells("A1:BH1")
    layouts["A1"] = "INTERNAL - schematic stock layouts (confirm kerf, grips and fabrication dimensions)"
    layouts["A1"].font = Font(size=13, bold=True, color=navy)
    for col in range(1, 61):
        layouts.column_dimensions[get_column_letter(col)].width = 2
    palette = ["007F86", "3671B7", "8064A2", "509D76"]
    for index, allocation in enumerate(allocations):
        row_index = 3 + index * 4
        layouts.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=60)
        layouts.cell(row_index, 1, _sheet_value(f"{allocation.get('row',{}).get('section','')} | {allocation.get('L',0)} m {allocation.get('kind','stock')} | {_money(allocation.get('cost'))}"))
        stock_mm = _numeric(allocation.get("L")) * 1000
        used, segments = 0, []
        for part in allocation.get("selected", []):
            used += _numeric(part.get("length"))
            segments.append(used)
        for col in range(1, 61):
            distance = (col - .5) / 60 * stock_mm
            piece = next((i for i, end in enumerate(segments) if distance <= end), None)
            layouts.cell(row_index+1, col).fill = PatternFill("solid", fgColor=palette[piece % len(palette)] if piece is not None else "D9E0E7")
        layouts.row_dimensions[row_index+1].height = 18
        layouts.merge_cells(start_row=row_index+2, start_column=1, end_row=row_index+2, end_column=60)
        cuts = " + ".join(f"{p.get('mark','')} {p.get('length',0)} mm" for p in allocation.get("selected", []))
        layouts.cell(row_index+2, 1, _sheet_value(cuts + (" | PROVISIONAL estimating tolerance" if allocation.get("provisional") else "")))
        layouts.cell(row_index+2, 1).alignment = Alignment(wrap_text=True)
        layouts.row_dimensions[row_index+2].height = 30
    layouts.sheet_properties.pageSetUpPr.fitToPage = True
    layouts.page_setup.orientation = "landscape"
    layouts.page_setup.paperSize = layouts.PAPERSIZE_A4
    layouts.page_setup.fitToWidth, layouts.page_setup.fitToHeight = 1, 0
    layouts.print_options.horizontalCentered = True
    layouts.print_title_rows = "1:2"
    plate_rows = [["Detailed plate" if state.get("plateMode") == "detailed" else "Inactive detailed plate", p.get("label", ""), p.get("qty", 0), p.get("length", 0), p.get("width", 0), p.get("thick", 0), _numeric(p.get("qty"))*_numeric(p.get("length"))*_numeric(p.get("width"))*_numeric(p.get("thick"))*density] for p in state.get("plates", [])]
    plate_rows.extend([["Purchased allocation", p.get("desc", ""), p.get("fraction", p.get("quantity", "")), "", "", p.get("thick", ""), p.get("weight", 0)] for p in result.get("plateSheets", [])])
    plate_rows.extend([["Net fabricated plate kg", "", "", "", "", "", result.get("plateWeight", 0)], ["Charged plate kg", "", "", "", "", "", result.get("plateGross", 0)], ["Plate buying cost AUD", "", "", "", "", "", result.get("plateBuy", 0)], ["Bolt internal cost AUD", "Unconfigured" if project.get("boltCost") in (None, "") else "Manual cost", "", "", "", "", project.get("boltCost", "")]])
    sheet("Plates and bolts", ["Type", "Description / mark", "Qty / sheets", "Length mm", "Width mm", "Thickness mm", "Weight kg / value"], plate_rows, [30, 65, 18, 18, 18, 18, 24])
    costs = [[i.get("name", ""), i.get("hours", ""), i.get("cost", 0), i.get("sell", 0)] for i in result.get("items", [])]
    costs.extend([["Installation", "", result.get("installCost", 0), result.get("instSell", 0)], ["FINAL TOTAL", "", result.get("totalCost", 0), result.get("totalSell", 0)]])
    ws = sheet("Costs and hours", ["Operation", "Hours", "Internal cost AUD", "Selling AUD"], costs, [50, 18, 25, 25])
    for row in ws.iter_rows(min_row=4, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = '"$"#,##0.00'
    notes = [["Review warning", str(w)] for w in result.get("warnings", [])]
    notes.extend(drawing_notes)
    notes += [["Project notes", project.get("notes", "")], ["Confirmed quotation scope", project.get("quoteScope", "")], ["Quotation inclusions / exclusions", project.get("exclusions", "")], ["Drawing / enquiry revision", project.get("revision", "")], ["Customer enquiry / scope", state.get("pdfScope", "")], ["Price basis", "AUD excluding GST; internal confidential estimate. Verify supplier availability, uncertain dimensions and missing quotes."]]
    for drawing in state.get("pdfFiles", []):
        notes.append(["Drawing reference", _text(drawing)])
    sheet("Notes and drawings", ["Type", "Details"], notes, [32, 115])
    sheet("Rate snapshot", ["Rate", "Applied value"], [[k, v] for k, v in sorted(result.get("effectiveRates", {}).items())], [55, 28])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def pdf_report(record: dict) -> bytes:
    state, result, project = _report_parts(record)
    drawing_notes = _drawing_review_notes(state)
    doc = fitz.open()
    page = None
    y = 0
    width, height = 595, 842
    left, right = 42, 553

    def ascii_text(value):
        return _text(value).replace("–", "-").replace("—", "-").replace("×", "x").replace("\t", " ").encode("latin-1", "replace").decode("latin-1")

    def new_page():
        nonlocal page, y
        page = doc.new_page(width=width, height=height)
        page.draw_rect(fitz.Rect(0, 0, width, 64), color=None, fill=(.09,.17,.28))
        page.insert_text((left, 30), "SteelTakeoff AI | INTERNAL ESTIMATE", fontsize=17, color=(1,1,1))
        page.insert_text((left, 48), "PRIVATE - NOT FOR CUSTOMER ISSUE | AUD excluding GST", fontsize=9, color=(.8,.87,.93))
        page.insert_text((left, height-24), f"Revision {record['revision']} | Page {len(doc)} | Verify unresolved takeoff and costing items", fontsize=8, color=(.4,.45,.5))
        y = 87

    def line(text="", size=10, bold=False, color=(.09,.17,.28), indent=0):
        nonlocal y
        text = ascii_text(text)
        for para in text.splitlines() or [""]:
            words = para.split(" ")
            lines, current = [], ""
            for word in words:
                # Split long evidence / filenames so no content silently runs off page.
                while fitz.get_text_length(word, fontname="helv", fontsize=size) > right-left-indent:
                    take = max(1, int((right-left-indent)/(size*.6)))
                    if current:
                        lines.append(current)
                        current = ""
                    lines.append(word[:take])
                    word = word[take:]
                candidate = f"{current} {word}" if current else word
                if fitz.get_text_length(candidate, fontname="helv", fontsize=size) > right-left-indent and current:
                    lines.append(current)
                    current = word
                else:
                    current = candidate
            lines.append(current)
            for text_line in lines:
                if page is None or y + size + 7 > height - 48:
                    new_page()
                page.insert_text((left+indent, y), text_line, fontsize=size, fontname="hebo" if bold else "helv", color=color)
                y += size + 5

    def heading(text):
        nonlocal y
        if page is None or y + 55 > height - 48:
            new_page()
        y += 12
        line(text, 12, True, (.0,.44,.49))

    new_page()
    line(project.get("name") or "Untitled project", 16, True)
    line(f"Customer: {project.get('client','')} | Address: {project.get('address','')}")
    line(f"Saved: {record.get('saved_at','')} | Scope: {project.get('scope','')}", 9)
    heading("Final estimate summary")
    line(f"Net steel including plates: {_numeric(result.get('net')):,.2f} kg | Purchased sections: {_numeric(result.get('purchase',{}).get('gross')):,.2f} kg")
    line(f"Internal known cost: {_money(result.get('totalCost'))} | Total sell: {_money(result.get('totalSell'))}", 12, True)
    margin = result.get("margin")
    line(f"Gross profit: {_money(result.get('profit'))} | Margin: {f'{margin:.2f}%' if isinstance(margin,(int,float)) else 'UNVERIFIED - incomplete costs'}")
    line(f"Supply/fabrication: {_money(result.get('supplySell'))} | Installation: {_money(result.get('instSell'))}")
    heading("Scope and assumptions")
    if project.get("quoteScope"):
        line("Confirmed quotation scope: " + _text(project["quoteScope"]))
    if project.get("exclusions"):
        line("Inclusions / exclusions: " + _text(project["exclusions"]))
    if project.get("revision"):
        line("Drawing / enquiry revision: " + _text(project["revision"]))
    line(project.get("notes") or "No project assumptions recorded.")
    if state.get("pdfScope"):
        line(state["pdfScope"])
    heading("Drawing references")
    for drawing in state.get("pdfFiles", []):
        line(_text(drawing), 9)
    if not state.get("pdfFiles"):
        line("No drawing references saved.")
    if drawing_notes:
        heading("Drawing analysis review - verify before pricing")
        for kind, detail in drawing_notes:
            line(kind, 9, True, color=(.65,.3,.05))
            line(detail, 9)
    heading("Steel material takeoff")
    for m in result.get("resolvedMembers", state.get("members", [])):
        kg = _numeric(m.get("qty"))*_numeric(m.get("length"))*_numeric(m.get("kgm"))/1000
        line(f"{m.get('mark','')} | {m.get('section','')} | {m.get('qty','')} x {m.get('length','')} mm | {m.get('kgm','')} kg/m | {kg:,.2f} kg | {m.get('finish','')}", 9, True)
        if m.get("sourceFile") or m.get("sourceEvidence"):
            line(f"Drawing: {m.get('sourceFile','')} page {m.get('sourcePage','')} | {m.get('sourceEvidence','')}", 8, indent=10)
    heading("Stock purchasing and cut layouts")
    for a in result.get("purchase", {}).get("allocations", []):
        if y + 60 > height - 48:
            new_page()
        row = a.get("row", {})
        line(f"{row.get('section','')} | {a.get('L',0)} m {a.get('kind','stock')} | Buy {_money(a.get('cost'))}", 9, True)
        cuts = a.get("selected", [])
        total = _numeric(a.get("L"))*1000
        if total > 0:
            page.draw_rect(fitz.Rect(left,y,right,y+13), color=(.65,.7,.75), fill=(.88,.9,.92))
            cursor = left
            for p in cuts:
                span = min(right-cursor,max(0,_numeric(p.get("length"))/total*(right-left)))
                if span > 0:
                    page.draw_rect(fitz.Rect(cursor,y,cursor+span,y+13), color=(1,1,1), fill=(.0,.5,.54))
                cursor += span
            y += 25
        line("Cuts: " + " + ".join(f"{p.get('mark','')} {p.get('length',0)} mm" for p in cuts), 8)
        if a.get("provisional"):
            line("PROVISIONAL saw tolerance - confirm measured lengths, stock overlength and kerf.", 8, color=(.65,.3,.05))
    heading("Plate and bolt allowances")
    line(f"Fabricated plates: {_numeric(result.get('plateWeight')):,.2f} kg | Charged sheets: {_numeric(result.get('plateGross')):,.2f} kg | Plate buy {_money(result.get('plateBuy'))}")
    for p in state.get("plates", []):
        active = "" if state.get("plateMode") == "detailed" else "INACTIVE detailed entry | "
        line(f"{active}{p.get('label','')} | {p.get('qty',0)} x {p.get('length',0)} x {p.get('width',0)} x {p.get('thick',0)} mm", 9)
    for p in result.get("plateSheets", []):
        line(f"{p.get('desc','')} | Equivalent sheets {p.get('fraction',p.get('quantity',''))} | {_numeric(p.get('weight')):,.2f} kg", 9)
    line(f"Bolt internal cost: {_money(project.get('boltCost')) if project.get('boltCost') not in (None,'') else 'UNCONFIGURED'}")
    heading("Hours, costs and selling")
    for item in result.get("items", []):
        hours = f" | {_numeric(item['hours']):,.2f} hours" if item.get("hours") is not None else ""
        line(f"{item.get('name','')}{hours} | Cost {_money(item.get('cost'))} | Sell {_money(item.get('sell'))}", 9)
    line(f"Installation | Cost {_money(result.get('installCost'))} | Sell {_money(result.get('instSell'))}", 9)
    heading("Unresolved takeoff and pricing issues")
    for warning in result.get("warnings", []):
        line(f"- {warning}", 9, color=(.65,.3,.05))
    if not result.get("warnings"):
        line("Review supplier availability, drawing completeness and fabrication dimensions before issue.")
    heading("Applied rate snapshot")
    for key, value in sorted(result.get("effectiveRates", {}).items()):
        line(f"{key}: {value}", 8)
    output = doc.tobytes(garbage=4, deflate=True)
    doc.close()
    return output


def create_router(data_dir: str | Path) -> APIRouter:
    router = APIRouter(prefix="/api")
    store = Store(data_dir)
    # Exposed for embedded/local tests without a global application singleton.
    router.store = store

    def error(exc):
        if isinstance(exc, RevisionConflict):
            raise HTTPException(409, str(exc)) from exc
        if isinstance(exc, KeyError):
            raise HTTPException(404, str(exc).strip("'")) from exc
        raise HTTPException(422, str(exc)) from exc

    @router.get("/projects")
    def list_projects():
        return {"projects": store.list_projects()}

    @router.post("/projects")
    def save_project(request: SaveRequest):
        try:
            result = _calculate(request.state, store)
            return store.save(request.state, result, request.id, request.expected_revision)
        except (ValueError, KeyError, TypeError) as exc:
            error(exc)

    @router.get("/projects/{project_id}")
    def get_project(project_id: str):
        try:
            return store.get(project_id)
        except KeyError as exc:
            error(exc)

    @router.get("/projects/{project_id}/revisions")
    def revisions(project_id: str):
        try:
            return {"revisions": store.revisions(project_id)}
        except KeyError as exc:
            error(exc)

    @router.get("/projects/{project_id}/revisions/{revision}")
    def get_revision(project_id: str, revision: int):
        try:
            return store.get(project_id, revision)
        except KeyError as exc:
            error(exc)

    @router.get("/settings")
    def settings():
        return store.settings()

    @router.put("/settings")
    def update_settings(request: dict):
        try:
            return store.update_settings(request)
        except ValueError as exc:
            error(exc)

    @router.post("/calculate")
    def calculate(request: StateRequest):
        try:
            return _calculate(request.state, store)
        except (ValueError, KeyError, TypeError) as exc:
            error(exc)

    @router.post("/email")
    def email(request: StateRequest):
        from estimating import build_email
        try:
            result = _calculate(request.state, store)
            return {"email": build_email(request.state, result)}
        except (ValueError, KeyError, TypeError) as exc:
            error(exc)

    @router.get("/projects/{project_id}/report.xlsx")
    def report_excel(project_id: str):
        try:
            record = store.get(project_id)
        except KeyError as exc:
            error(exc)
        return Response(excel_report(record), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{_safe_name(record["name"], "_INTERNAL.xlsx")}"'})

    @router.get("/projects/{project_id}/report.pdf")
    def report_pdf(project_id: str):
        try:
            record = store.get(project_id)
        except KeyError as exc:
            error(exc)
        return Response(pdf_report(record), media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{_safe_name(record["name"], "_INTERNAL.pdf")}"'})

    @router.get("/backup")
    def backup():
        return Response(store.backup(), media_type="application/zip", headers={"Content-Disposition": f'attachment; filename="SteelTakeoff_backup_{datetime.now().strftime("%Y%m%d_%H%M%S")}.zip"'})

    @router.post("/restore")
    async def restore(file: UploadFile = File(...)):
        payload = await file.read(Store.MAX_ARCHIVE_BYTES + 1)
        try:
            return store.restore(payload)
        except (ValueError, OSError, TypeError) as exc:
            error(exc)

    @router.post("/sections/import")
    async def sections_import(file: UploadFile = File(...), verified: bool = Form(False), source: str = Form("")):
        payload = await file.read(15 * 1024 * 1024 + 1)
        try:
            imported = import_sections(payload, file.filename or "", verified, source)
            with store.lock:
                merged = {normal_section(s["section"]): s for s in store.settings()["sections"]}
                merged.update({normal_section(s["section"]): s for s in imported})
                result = store.update_settings({"sections": list(merged.values())})
            return {"imported": len(imported), **result}
        except (ValueError, UnicodeDecodeError, OSError) as exc:
            error(exc)

    return router
