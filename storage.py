"""Local SQLite project history and editable, verified estimating settings.

Each connection is short lived. Project JSON is preserved verbatim so older V2
exports and future fields survive a save/load round trip.
"""
from __future__ import annotations

import copy
import csv
import io
import json
import math
import os
import re
import shutil
import sqlite3
import stat
import tempfile
import threading
import uuid
import zipfile
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any


class RevisionConflict(ValueError):
    pass


def utcnow() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="microseconds")


def json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(",", ":"))


def number(value: Any, name: str, minimum: float = 0, maximum: float = 1e9) -> float:
    if isinstance(value, bool):
        raise ValueError(f"{name} must be a number")
    try:
        val = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"{name} must be a number") from None
    if not math.isfinite(val) or not minimum <= val <= maximum:
        raise ValueError(f"{name} must be between {minimum:g} and {maximum:g}")
    return val


def normal_section(section: str) -> str:
    return re.sub(r"\s+", "", str(section)).replace("×", "X").upper()


FAMILIES = {"UB", "UC", "PFC", "SHS", "RHS", "CHS", "Pipe", "Angle", "Flat bar"}


def validate_section(raw: dict) -> dict:
    if not isinstance(raw, dict):
        raise ValueError("Each section must be an object")
    row = copy.deepcopy(raw)
    row["section"] = str(raw.get("section", raw.get("designation", ""))).strip()
    if not row["section"] or len(row["section"]) > 100:
        raise ValueError("Section designation is required (maximum 100 characters)")
    family = str(raw.get("family", "")).strip()
    family = {"PIPE": "Pipe", "PIPES": "Pipe", "ANGLES": "Angle", "ANGLE": "Angle", "FLAT BAR": "Flat bar", "FLAT": "Flat bar"}.get(family.upper(), family.upper())
    if family not in FAMILIES:
        raise ValueError(f"Unsupported section family: {family}")
    row["family"] = family
    row["kgm"] = number(raw.get("kgm"), "Section kg/m", .000001, 10000)
    verified = raw.get("verified") is True or str(raw.get("verified", "")).strip().lower() in ("true", "yes", "1", "verified")
    if not verified:
        raise ValueError(f"{row['section']}: confirm that its kg/m comes from a verified source")
    row["verified"] = True
    row["source"] = str(raw.get("source", "")).strip()
    if not row["source"]:
        raise ValueError(f"{row['section']}: a verification source is required")
    stocks = raw.get("stocks", raw.get("stock_lengths", []))
    if isinstance(stocks, str):
        stocks = re.split(r"[,;|\s]+", stocks.strip()) if stocks.strip() else []
    if not isinstance(stocks, list) or len(stocks) > 50:
        raise ValueError("Stock lengths must be a list of lengths in metres")
    row["stocks"] = sorted({number(x, "Stock length (m)", .001, 100) for x in stocks})
    supplier_rate = raw.get("supplier_rate", raw.get("supplier_price"))
    if supplier_rate not in (None, ""):
        row["supplier_rate"] = number(supplier_rate, "Supplier rate ($/tonne)", 0, 1000000)
    else:
        row["supplier_rate"] = None
    # UB and UC mass is encoded by the designation; reject a conflicting table.
    mass = re.fullmatch(r"\d+(?:\.\d+)?(?:UB|UC)(\d+(?:\.\d+)?)", normal_section(row["section"]))
    if mass and abs(float(mass[1]) - row["kgm"]) > .051:
        raise ValueError(f"{row['section']}: kg/m disagrees with designation")
    row.update(designation=row["section"], stock_lengths=row["stocks"], supplier_price=row["supplier_rate"])
    return row


def default_settings() -> dict:
    from estimating import DEFAULT_RATES
    return {"rates": copy.deepcopy(DEFAULT_RATES), "sections": [], "assemblies": [
        {"name": "100x100x5 SHS post with top/bottom plates", "hours": 1.0, "source": "User estimating allowance"},
        {"name": "50x50x5 SHS post with top/bottom plates", "hours": .5, "source": "User estimating allowance"},
    ]}


def validate_settings(incoming: dict, current: dict) -> dict:
    if not isinstance(incoming, dict):
        raise ValueError("Settings must be an object")
    result = copy.deepcopy(current)
    if "rates" in incoming:
        rates = incoming["rates"]
        if not isinstance(rates, dict):
            raise ValueError("Rates must be an object")
        for key, value in rates.items():
            if key not in result["rates"]:
                raise ValueError(f"Unknown rate: {key}")
            minimum = .000000000001 if key in {"plate_reference_weight", "saw_sell_per_hour", "steel_density_kg_mm3"} else 0
            validated = number(value, key, minimum, 1e9)
            if "margin" in key.lower() and validated >= 100:
                raise ValueError(f"{key}: gross margin must be below 100%")
            result["rates"][key] = validated
    if "sections" in incoming:
        if not isinstance(incoming["sections"], list) or len(incoming["sections"]) > 20000:
            raise ValueError("Sections must be a list with at most 20,000 entries")
        result["sections"] = [validate_section(row) for row in incoming["sections"]]
        keys = [normal_section(row["section"]) for row in result["sections"]]
        if len(set(keys)) != len(keys):
            raise ValueError("Duplicate section designation")
    if "assemblies" in incoming:
        rows = incoming["assemblies"]
        if not isinstance(rows, list) or len(rows) > 2000:
            raise ValueError("Assemblies must be a list with at most 2,000 entries")
        result["assemblies"] = []
        for row in rows:
            if not isinstance(row, dict) or not str(row.get("name", "")).strip():
                raise ValueError("Each assembly requires a name")
            clean = copy.deepcopy(row)
            clean["hours"] = number(row.get("hours", row.get("labour_hours")), "Assembly hours", 0, 10000)
            clean["labour_hours"] = clean["hours"]
            if "efficiency" in row and "repeat_efficiency" not in row:
                clean["repeat_efficiency"] = number(row["efficiency"], "Repeat efficiency", 0, 1) * 100
            if "repeat_efficiency" in row:
                clean["repeat_efficiency"] = number(row["repeat_efficiency"], "Repeat efficiency (%)", 0, 100)
            if "repeat_efficiency" in clean:
                clean["efficiency"] = clean["repeat_efficiency"] / 100
            result["assemblies"].append(clean)
    json_text(result)
    return result


class Store:
    DB_NAME = "steeltakeoff.sqlite3"
    MAX_ARCHIVE_BYTES = 150 * 1024 * 1024
    MAX_EXPANDED_BYTES = 400 * 1024 * 1024

    def __init__(self, data_dir: str | Path):
        self.data_dir = Path(data_dir).resolve()
        self.lock = threading.RLock()
        self.data_dir.mkdir(parents=True, exist_ok=True)
        self._initialise()

    @property
    def db_path(self) -> Path:
        return self.data_dir / self.DB_NAME

    @contextmanager
    def connect(self):
        db = sqlite3.connect(self.db_path, timeout=15)
        db.row_factory = sqlite3.Row
        db.execute("PRAGMA foreign_keys=ON")
        try:
            with db:
                yield db
        finally:
            db.close()

    def _initialise(self):
        with self.connect() as db:
            db.executescript("""
            CREATE TABLE IF NOT EXISTS projects (
                id TEXT PRIMARY KEY, name TEXT NOT NULL, client TEXT NOT NULL,
                revision INTEGER NOT NULL, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS revisions (
                project_id TEXT NOT NULL REFERENCES projects(id), revision INTEGER NOT NULL,
                saved_at TEXT NOT NULL, state_json TEXT NOT NULL, result_json TEXT NOT NULL,
                PRIMARY KEY(project_id, revision)
            );
            CREATE TABLE IF NOT EXISTS settings (id INTEGER PRIMARY KEY CHECK(id=1), data_json TEXT NOT NULL);
            PRAGMA user_version=1;
            """)
            db.execute("INSERT OR IGNORE INTO settings(id,data_json) VALUES(1,?)", (json_text(default_settings()),))

    def settings(self) -> dict:
        with self.lock, self.connect() as db:
            data = json.loads(db.execute("SELECT data_json FROM settings WHERE id=1").fetchone()[0])
        defaults = default_settings()
        defaults.update(data)
        defaults["rates"] = {**default_settings()["rates"], **data.get("rates", {})}
        return defaults

    def update_settings(self, incoming: dict) -> dict:
        with self.lock:
            clean = validate_settings(incoming, self.settings())
            with self.connect() as db:
                db.execute("UPDATE settings SET data_json=? WHERE id=1", (json_text(clean),))
        return clean

    def list_projects(self) -> list[dict]:
        with self.lock, self.connect() as db:
            return [dict(r) for r in db.execute("SELECT * FROM projects ORDER BY updated_at DESC")]

    def save(self, state: dict, result: dict, project_id: str | None = None, expected_revision: int | None = None) -> dict:
        if not isinstance(state, dict) or not isinstance(state.get("project"), dict):
            raise ValueError("Project state must contain a project object")
        if not isinstance(state.get("members", []), list) or not isinstance(state.get("plates", []), list):
            raise ValueError("Members and plates must be arrays")
        state_json = json_text(state)
        if len(state_json.encode("utf-8")) > 20 * 1024 * 1024:
            raise ValueError("Project JSON exceeds 20 MB")
        result_json = json_text(result)
        project_id = project_id or uuid.uuid4().hex
        if not re.fullmatch(r"[A-Za-z0-9_-]{1,80}", project_id):
            raise ValueError("Invalid project ID")
        with self.lock, self.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            old = db.execute("SELECT * FROM projects WHERE id=?", (project_id,)).fetchone()
            revision = old["revision"] if old else 0
            if expected_revision is not None and expected_revision != revision:
                raise RevisionConflict(f"Project changed: expected revision {expected_revision}, current revision is {revision}. Reload before saving.")
            now = utcnow()
            revision += 1
            project = state["project"]
            created_at = old["created_at"] if old else now
            db.execute("""INSERT INTO projects VALUES(?,?,?,?,?,?)
                ON CONFLICT(id) DO UPDATE SET name=excluded.name,client=excluded.client,
                revision=excluded.revision,updated_at=excluded.updated_at""",
                (project_id, str(project.get("name", "")), str(project.get("client", "")), revision, created_at, now))
            db.execute("INSERT INTO revisions VALUES(?,?,?,?,?)", (project_id, revision, now, state_json, result_json))
        # A subsequent save can commit immediately after this transaction. Return
        # this committed revision rather than querying the mutable latest record.
        return {"id": project_id, "name": str(project.get("name", "")), "client": str(project.get("client", "")),
                "revision": revision, "created_at": created_at, "updated_at": now, "saved_at": now,
                "state": json.loads(state_json), "result": json.loads(result_json)}

    def get(self, project_id: str, revision: int | None = None) -> dict:
        with self.lock, self.connect() as db:
            info = db.execute("SELECT * FROM projects WHERE id=?", (project_id,)).fetchone()
            if info is None:
                raise KeyError("Project not found")
            row = db.execute("SELECT * FROM revisions WHERE project_id=? AND revision=?", (project_id, revision if revision is not None else info["revision"])).fetchone()
            if row is None:
                raise KeyError("Project revision not found")
            return {**dict(info), "revision": row["revision"], "saved_at": row["saved_at"], "state": json.loads(row["state_json"]), "result": json.loads(row["result_json"])}

    def revisions(self, project_id: str) -> list[dict]:
        self.get(project_id)
        with self.lock, self.connect() as db:
            return [dict(r) for r in db.execute("SELECT revision,saved_at FROM revisions WHERE project_id=? ORDER BY revision DESC", (project_id,))]

    def backup(self) -> bytes:
        """SQLite online backup captures committed records even with an active WAL."""
        output = io.BytesIO()
        with self.lock, tempfile.TemporaryDirectory(prefix="steel-backup-") as temp:
            backup_db = Path(temp) / self.DB_NAME
            dest = sqlite3.connect(backup_db)
            try:
                with self.connect() as source:
                    source.backup(dest)
            finally:
                dest.close()
            with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("manifest.json", json_text({"application": "SteelTakeoff AI", "format": 1, "created_at": utcnow(), "database": self.DB_NAME}))
                archive.write(backup_db, self.DB_NAME)
                uploads = self.data_dir / "uploads"
                if uploads.is_dir():
                    for path in sorted(uploads.rglob("*")):
                        if path.is_file() and not path.is_symlink() and path.resolve().is_relative_to(uploads.resolve()):
                            archive.write(path, path.relative_to(self.data_dir).as_posix())
        return output.getvalue()

    def restore(self, payload: bytes) -> dict:
        """Validate in a staging directory, then preserve and replace local data."""
        if len(payload) > self.MAX_ARCHIVE_BYTES:
            raise ValueError("Backup archive exceeds 150 MB")
        stage = Path(tempfile.mkdtemp(prefix=f".{self.data_dir.name}-restore-", dir=self.data_dir.parent))
        previous = self.data_dir.parent / f".{self.data_dir.name}-before-restore-{uuid.uuid4().hex[:12]}"
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                names = set()
                expanded = 0
                for entry in archive.infolist():
                    name = entry.filename
                    part = PurePosixPath(name)
                    if (name in names or "\\" in name or part.is_absolute() or ".." in part.parts
                            or ":" in name or not part.parts or stat.S_ISLNK(entry.external_attr >> 16)):
                        raise ValueError("Unsafe or duplicate backup archive path")
                    names.add(name)
                    if name not in {"manifest.json", self.DB_NAME} and part.parts[0] != "uploads":
                        raise ValueError("Backup contains unsupported files")
                    expanded += entry.file_size
                    if expanded > self.MAX_EXPANDED_BYTES or len(names) > 10000:
                        raise ValueError("Backup expands beyond permitted size")
                if not {"manifest.json", self.DB_NAME}.issubset(names):
                    raise ValueError("Not a SteelTakeoff backup: missing manifest or database")
                manifest = json.loads(archive.read("manifest.json"))
                if not isinstance(manifest, dict) or manifest.get("application") != "SteelTakeoff AI" or manifest.get("format") != 1:
                    raise ValueError("Unsupported backup format")
                for entry in archive.infolist():
                    if entry.filename == "manifest.json":
                        continue
                    dest = stage.joinpath(*PurePosixPath(entry.filename).parts)
                    if entry.is_dir():
                        dest.mkdir(parents=True, exist_ok=True)
                    else:
                        dest.parent.mkdir(parents=True, exist_ok=True)
                        with archive.open(entry) as src, dest.open("wb") as target:
                            shutil.copyfileobj(src, target)
            # Never run arbitrary schema or migration code from an uploaded DB.
            self._validate_backup_database(stage / self.DB_NAME)
            (stage / "uploads").mkdir(exist_ok=True)
            with self.lock:
                os.replace(self.data_dir, previous)
                try:
                    os.replace(stage, self.data_dir)
                except BaseException:
                    os.replace(previous, self.data_dir)
                    raise
            return {"restored": True, "preserved_previous_data": str(previous), "projects": len(self.list_projects())}
        except (zipfile.BadZipFile, json.JSONDecodeError, sqlite3.DatabaseError, RuntimeError) as exc:
            raise ValueError(f"Invalid backup: {exc}") from None
        finally:
            if stage.exists():
                shutil.rmtree(stage)

    @staticmethod
    def _validate_backup_database(path: Path):
        db = sqlite3.connect(f"{path.as_uri()}?mode=ro", uri=True)
        try:
            if db.execute("PRAGMA quick_check").fetchone()[0] != "ok":
                raise ValueError("Backup database integrity check failed")
            schema = db.execute("SELECT type,name,sql FROM sqlite_master WHERE name NOT LIKE 'sqlite_%'").fetchall()
            if {row[1] for row in schema if row[0] == "table"} != {"projects", "revisions", "settings"} or any(row[0] not in {"table", "index"} for row in schema):
                raise ValueError("Backup database schema is unsupported")
            expected = {"projects": ["id", "name", "client", "revision", "created_at", "updated_at"], "revisions": ["project_id", "revision", "saved_at", "state_json", "result_json"], "settings": ["id", "data_json"]}
            for table, columns in expected.items():
                if [r[1] for r in db.execute(f"PRAGMA table_info({table})")] != columns:
                    raise ValueError("Backup database columns are unsupported")
            setting = db.execute("SELECT data_json FROM settings WHERE id=1").fetchone()
            if not setting:
                raise ValueError("Backup settings missing")
            validate_settings(json.loads(setting[0]), default_settings())
            for state, result in db.execute("SELECT state_json,result_json FROM revisions"):
                data = json.loads(state)
                if not isinstance(data, dict) or not isinstance(data.get("project"), dict) or not isinstance(json.loads(result), dict):
                    raise ValueError("Backup contains invalid project data")
                json_text(data)
                json_text(json.loads(result))
            if db.execute("PRAGMA foreign_key_check").fetchone():
                raise ValueError("Backup project history is inconsistent")
            if db.execute("SELECT 1 FROM projects p LEFT JOIN revisions r ON p.id=r.project_id AND p.revision=r.revision WHERE r.project_id IS NULL").fetchone():
                raise ValueError("Backup project revision is missing")
        finally:
            db.close()


def import_sections(payload: bytes, filename: str, verified: bool = False, source: str = "") -> list[dict]:
    if len(payload) > 15 * 1024 * 1024:
        raise ValueError("Section table exceeds 15 MB")
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        reader = csv.DictReader(io.StringIO(payload.decode("utf-8-sig")))
        raw_rows = list(reader)
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
        from xml.etree.ElementTree import ParseError
        # Reject oversized decompressed spreadsheets before the XML parser.
        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                if sum(r.file_size for r in archive.infolist()) > 80 * 1024 * 1024:
                    raise ValueError("Excel table expands beyond permitted size")
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        except (zipfile.BadZipFile, KeyError, InvalidFileException, ParseError, OSError) as exc:
            raise ValueError(f"Invalid Excel section table: {exc}") from None
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = next(rows, None)
            if headers is None:
                raise ValueError("Section table is empty")
            raw_rows = []
            for row in rows:
                if len(raw_rows) >= 20000:
                    raise ValueError("Section table exceeds 20,000 rows")
                if any(x is not None for x in row):
                    raw_rows.append(dict(zip(headers, row)))
        finally:
            workbook.close()
    else:
        raise ValueError("Upload a UTF-8 .csv or .xlsx verified section table")
    if not raw_rows or len(raw_rows) > 20000:
        raise ValueError("Section table must contain 1 to 20,000 rows")
    aliases = {"designation": "section", "section_designation": "section", "kg/m": "kgm", "kg_per_m": "kgm", "weight_kg_m": "kgm", "stock_lengths": "stocks", "stock_lengths_m": "stocks", "supplier_rate_per_tonne": "supplier_rate", "verification_source": "source"}
    clean = []
    for i, raw in enumerate(raw_rows, 2):
        normal = {}
        for key, val in raw.items():
            if key is None:
                continue
            name = str(key).strip().lower().replace(" ", "_")
            normal[aliases.get(name, name)] = val
        normal["verified"] = verified or normal.get("verified", False)
        normal["source"] = normal.get("source") or source
        try:
            clean.append(validate_section(normal))
        except ValueError as exc:
            raise ValueError(f"Row {i}: {exc}") from None
    keys = [normal_section(r["section"]) for r in clean]
    if len(keys) != len(set(keys)):
        raise ValueError("Imported table contains duplicate section designations")
    return clean
