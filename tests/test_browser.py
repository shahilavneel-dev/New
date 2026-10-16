"""Real Chromium workflow checks against an isolated local server/database."""
import os
from pathlib import Path
import shutil
import socket
import subprocess
import sys
import time

import httpx
import pytest

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope='module')
def live_app(tmp_path_factory):
    with socket.socket() as sock:
        sock.bind(('127.0.0.1', 0))
        port = sock.getsockname()[1]
    data = tmp_path_factory.mktemp('browser-data')
    process = subprocess.Popen(
        [sys.executable, '-m', 'uvicorn', 'server:app', '--host', '127.0.0.1', '--port', str(port), '--log-level', 'warning'],
        cwd=ROOT, env={**os.environ, 'STEELTAKEOFF_DATA_DIR': str(data), 'STEELTAKEOFF_NO_BROWSER': '1'},
        stdout=subprocess.DEVNULL, stderr=subprocess.PIPE,
    )
    base = f'http://127.0.0.1:{port}'
    try:
        for _ in range(100):
            if process.poll() is not None:
                pytest.fail(process.stderr.read().decode())
            try:
                if httpx.get(base+'/api/health').status_code == 200:
                    break
            except httpx.RequestError:
                pass
            time.sleep(.1)
        else:
            pytest.fail('Local browser test server did not become ready')
        yield base
    finally:
        process.terminate()
        try:
            process.wait(timeout=8)
        except subprocess.TimeoutExpired:
            process.kill()
            process.wait(timeout=3)
        process.stderr.close()


@pytest.fixture
def browser_page(live_app):
    playwright = pytest.importorskip('playwright.sync_api')
    with playwright.sync_playwright() as p:
        executable = shutil.which('chromium') or shutil.which('google-chrome')
        browser = p.chromium.launch(executable_path=executable, headless=True, args=['--no-sandbox'])
        page = browser.new_page(viewport={'width':1700,'height':1050})
        errors = []
        page.on('pageerror',lambda error: errors.append(str(error)))
        page.on('dialog',lambda dialog: dialog.accept('1000') if dialog.type=='prompt' else dialog.accept())
        page.goto(live_app)
        page.get_by_role('button',name='Overview & projects',exact=True).wait_for()
        yield page
        browser.close()
        assert errors == []


def navigate(page, label, heading):
    page.get_by_role('button',name=label,exact=True).click()
    page.get_by_role('heading',name=heading,exact=True).wait_for()


def test_estimate_save_reopen_reports_email(browser_page,live_app,tmp_path):
    page=browser_page
    page.get_by_role('button',name='Load labelled example',exact=True).click()
    page.get_by_role('heading',name='Job setup').wait_for()
    field=page.locator("input[onchange*=\"project.name\"]")
    field.fill('Browser integration estimate')
    field.press('Tab')
    page.wait_for_function("state.project.name==='Browser integration estimate'")
    page.evaluate('persistProject(false)')
    saved=page.evaluate('projectMeta')
    assert saved['revision']>=1
    record=httpx.get(live_app+'/api/projects/'+saved['id']).json()
    assert record['state']['project']['name']=='Browser integration estimate'
    assert record['result']['net']>0
    navigate(page,'Steel stock nesting','Stock selection and nesting')
    assert page.locator('.bar').count()>0
    navigate(page,'Reports & quotation','Customer quotation email')
    email=page.locator('#emailText').input_value()
    assert 'Supply & Fabrication: $' in email
    assert all(token not in email.lower() for token in ['internal','gross profit','supplier buy','70.01','60.68'])
    page.locator('#emailText').fill(email+'\nCustom lead time: six weeks.')
    navigate(page,'Enquiry & scope','Job setup')
    navigate(page,'Reports & quotation','Customer quotation email')
    assert 'Custom lead time: six weeks.' in page.locator('#emailText').input_value()
    with page.expect_download() as info:
        page.get_by_role('button',name='Download Excel',exact=True).click()
    path=tmp_path/'estimate.xlsx'
    info.value.save_as(path)
    from openpyxl import load_workbook
    book=load_workbook(path)
    assert 'Members' in book.sheetnames and 'Stock layouts' in book.sheetnames
    with page.expect_download() as info:
        page.get_by_role('button',name='Download PDF',exact=True).click()
    path=tmp_path/'estimate.pdf'
    info.value.save_as(path)
    import fitz
    with fitz.open(path) as doc:
        assert 'INTERNAL ESTIMATE' in doc[0].get_text()
    navigate(page,'Overview & projects','Your estimates')
    card=page.locator('.project-card').filter(has_text='Browser integration estimate')
    card.get_by_role('button',name='Open estimate').click()
    page.get_by_role('heading',name='Job setup').wait_for()
    assert page.locator("input[onchange*=\"project.name\"]").input_value()=='Browser integration estimate'
    # An omitted required bolt cost must suppress a verified percentage.
    page.evaluate("state.project.boltCost='';save();render()")
    page.wait_for_function('calculationCache.cost_complete===false')
    assert 'Incomplete' in page.locator('#metrics').inner_text()


def test_sample_pdf_review_gate_highlight_and_measurement(browser_page,live_app):
    page=browser_page
    navigate(page,'Drawings & AI takeoff','Drawing review & assisted takeoff')
    page.locator('input[type=file][multiple]').set_input_files(str(ROOT/'SAMPLE_TEST_DRAWING.pdf'))
    page.get_by_role('heading',name='Candidate takeoff — 3 reference(s)').wait_for()
    candidates=page.evaluate('state.pdfCandidates')
    assert [(c['mark'],c['qty'],c['length_mm']) for c in candidates]==[('B1',2,4500),('C1',4,3000),('B2',1,5500)]
    assert candidates[1]['kgm'] is None
    assert page.get_by_role('button',name='Import 0 approved line(s)').is_disabled()
    page.get_by_role('button',name='Page 1',exact=True).first.click()
    page.locator('.pdfHighlight').wait_for()
    page.get_by_role('button',name='+',exact=True).click()
    page.wait_for_function('pdfZoom===1.25')
    page.get_by_role('button',name='Fit width',exact=True).click()
    page.get_by_role('button',name='Approve',exact=True).nth(1).click()
    assert page.evaluate("state.pdfCandidates[1].status")=='pending'
    page.get_by_role('button',name='Approve',exact=True).first.click()
    page.get_by_role('button',name='Import 1 approved line(s)').wait_for()
    page.get_by_role('button',name='Import 1 approved line(s)').click()
    page.get_by_role('heading',name='Material takeoff',exact=True).wait_for()
    assert page.evaluate('state.members.length')==1
    member=page.evaluate('state.members[0]')
    assert member['mark']=='B1' and member['sourcePage']==1 and member['sourceBBox']
    page.evaluate('persistProject(false)')
    navigate(page,'Drawings & AI takeoff','Drawing review & assisted takeoff')
    page.get_by_role('button',name='Calibrate',exact=True).click()
    # Provide an explicit known reference through the actual calibration dialog.
    page.locator('#pdfDrawingImage').wait_for()
    page.locator('.pdfDrawingStage').click(position={'x':40,'y':140})
    page.wait_for_function('pdfPoints.length===1')
    page.locator('.pdfDrawingStage').click(position={'x':140,'y':140})
    page.wait_for_function('Object.keys(state.pdfCalibrations).length===1')
    page.get_by_role('button',name='Measure',exact=True).click()
    page.locator('.pdfDrawingStage').click(position={'x':40,'y':170})
    page.wait_for_function('pdfPoints.length===1')
    page.locator('.pdfDrawingStage').click(position={'x':140,'y':170})
    page.wait_for_function('state.pdfMeasurements.length===1')
    assert page.evaluate('state.pdfMeasurements[0].provisional') is True
    assert abs(page.evaluate('state.pdfMeasurements[0].length_mm')-1000)<=2
    page.evaluate('persistProject(false)')
    assert not httpx.get(live_app+'/api/health').json()['ai_ready']


def test_cross_file_level_evidence_stays_provisional(browser_page,tmp_path):
    import fitz
    paths=[]
    for name,text in [('structural.pdf','DRAWING NO S101 REV A\nC1 100x100x5 SHS QTY 1\nC1 BASE RL 9.500'),('architectural.pdf','DRAWING NO A101 REV A\nC1 TOP RL 12.500')]:
        path=tmp_path/name
        with fitz.open() as doc:
            doc.new_page().insert_text((40,40),text,fontsize=12)
            doc.save(path)
        paths.append(str(path))
    page=browser_page
    navigate(page,'Drawings & AI takeoff','Drawing review & assisted takeoff')
    page.locator('input[type=file][multiple]').set_input_files(paths)
    page.wait_for_function('state.pdfFiles.length===2 && pdfBusy===false')
    page.get_by_role('button',name='Cross-reference local drawings',exact=True).click()
    page.wait_for_function('state.pdfCrossReferences?.rl_suggestions?.length===1')
    result=page.evaluate('state.pdfCrossReferences')
    assert result['rl_suggestions'][0]['suggested_length_mm']==3000
    assert len({ref['file_id'] for ref in result['rl_suggestions'][0]['references']})==2
    assert page.evaluate('state.members.length')==0
    assert page.evaluate('state.pdfCandidates[0].length_mm') is None
    assert 'Provisional' in page.locator('#view').inner_text()
    page.evaluate('persistProject(false)')
