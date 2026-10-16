import copy
import io
import json
import zipfile
from concurrent.futures import ThreadPoolExecutor

import fitz
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from project_api import create_router
from storage import RevisionConflict, Store, import_sections


@pytest.fixture
def state():
    return {
        "project": {"name": "Test job", "client": "Customer", "scope": "sf", "notes": "Confirm site dimensions", "boltEnabled": False, "delivery": "none", "kerf": 3, "sawTol": 10},
        "members": [{"mark": "B1", "family": "UB", "section": "250UB31.4", "kgm": 31.4, "qty": 2, "length": 4000, "finish": "Raw", "process": "saw", "cuts": 1, "fabEach": 1, "efficiency": 100}],
        "plates": [], "plateMode": "none", "draftMode": "none", "fabHoursT": 18,
        "pdfFiles": [{"id": "drawing-1", "name": "A101.pdf", "revision": "B"}],
        "future_field": {"preserve": [1, "legacy V2", {"anything": True}]},
    }


@pytest.fixture
def client(tmp_path):
    app = FastAPI()
    app.include_router(create_router(tmp_path / "app_data"))
    return TestClient(app)


def test_revision_roundtrip_conflict_and_history(tmp_path, state):
    store = Store(tmp_path / "app_data")
    first = store.save(state, {"totalSell": 10}, expected_revision=0)
    assert first["revision"] == 1
    assert first["state"] == state
    next_state = copy.deepcopy(state)
    next_state["project"]["name"] = "Revised job"
    second = store.save(next_state, {"totalSell": 20}, first["id"], 1)
    assert second["revision"] == 2
    assert store.get(first["id"], 1)["state"] == state
    with pytest.raises(RevisionConflict):
        store.save(state, {}, first["id"], 1)
    assert [r["revision"] for r in store.revisions(first["id"])] == [2, 1]
    reloaded = Store(tmp_path / "app_data")
    assert reloaded.get(first["id"])["state"] == next_state


def test_concurrent_save_responses_identify_their_own_committed_revision(tmp_path):
    store = Store(tmp_path / "app_data")

    def save(index):
        state = {"project": {"name": f"Writer {index}"}, "members": [], "plates": []}
        result = store.save(state, {"writer": index}, "concurrent")
        assert result["state"] == state
        assert result["name"] == state["project"]["name"]
        assert result["result"]["writer"] == index
        return result["revision"]

    with ThreadPoolExecutor(max_workers=12) as pool:
        revisions = list(pool.map(save, range(50)))
    assert sorted(revisions) == list(range(1, 51))


def test_api_revision_and_missing_projects(client, state):
    saved = client.post("/api/projects", json={"state": state, "expected_revision": 0})
    assert saved.status_code == 200, saved.text
    record = saved.json()
    pid = record["id"]
    assert record["state"]["future_field"] == state["future_field"]
    assert client.post("/api/projects", json={"id": pid, "state": state, "expected_revision": 0}).status_code == 409
    assert client.get(f"/api/projects/{pid}/revisions").json()["revisions"][0]["revision"] == 1
    assert client.get(f"/api/projects/{pid}/revisions/1").json()["state"] == state
    assert client.get("/api/projects/missing").status_code == 404
    assert client.get(f"/api/projects/{pid}/revisions/99").status_code == 404
    assert client.post("/api/projects", json={"state": {"project": {}, "members": "bad"}}).status_code == 422
    assert client.post("/api/projects", json={"state": state, "expected_revision": True}).status_code == 422


def test_settings_partial_validated_and_no_weights_invented(client):
    original = client.get("/api/settings").json()
    assert original["sections"] == []
    changed = client.put("/api/settings", json={"rates": {"structural_stock_per_tonne": 2100}})
    assert changed.status_code == 200
    assert changed.json()["rates"]["structural_stock_per_tonne"] == 2100
    assert changed.json()["rates"]["saw_heavy_per_cut"] == original["rates"]["saw_heavy_per_cut"]
    for invalid in ({"rates": {"not_a_rate": 1}}, {"rates": {"material_sf_margin": 100}}, {"rates": {"plate_reference_weight": 0}}):
        assert client.put("/api/settings", json=invalid).status_code == 422
    entry = {"designation": "100x100x5 SHS", "family": "SHS", "kgm": 14.9, "stock_lengths": [8, 12], "supplier_price": 2200, "verified": True, "source": "Supplier verified catalogue"}
    changed = client.put("/api/settings", json={"sections": [entry]})
    assert changed.status_code == 200, changed.text
    row = changed.json()["sections"][0]
    assert row["section"] == row["designation"] == entry["designation"]
    assert row["stocks"] == [8, 12]
    assert row["supplier_rate"] == 2200
    assert client.put("/api/settings", json={"sections": [{**entry, "source": ""}]}).status_code == 422
    assert client.put("/api/settings", json={"sections": [{**entry, "kgm": -1}]}).status_code == 422
    assert client.put("/api/settings", json={"sections": [{**entry, "stock_lengths": [8000]}]}).status_code == 422
    mismatch = {**entry, "designation": "250UB31.4", "family": "UB", "kgm": 34.5}
    assert client.put("/api/settings", json={"sections": [mismatch]}).status_code == 422


def test_rate_snapshot_and_overrides_are_reproducible(client, state):
    client.put("/api/settings", json={"rates": {"structural_stock_per_tonne": 2200}})
    state["rates"] = {"structural_stock_per_tonne": 2300}
    saved = client.post("/api/projects", json={"state": state}).json()
    assert saved["result"]["effectiveRates"]["structural_stock_per_tonne"] == 2300
    client.put("/api/settings", json={"rates": {"structural_stock_per_tonne": 2400}})
    assert client.get(f"/api/projects/{saved['id']}").json()["result"]["effectiveRates"]["structural_stock_per_tonne"] == 2300
    state.pop("rates")
    new = client.post("/api/calculate", json={"state": state}).json()
    assert new["effectiveRates"]["structural_stock_per_tonne"] == 2400


def test_catalogue_fills_missing_mass_but_preserves_invalid_zero_for_review(client, state):
    entry = {"section": "250UB31.4", "family": "UB", "kgm": 31.4, "stocks": [9], "verified": True, "source": "Published designation mass"}
    assert client.put("/api/settings", json={"sections": [entry]}).status_code == 200
    state["members"][0]["kgm"] = ""
    inferred = client.post("/api/calculate", json={"state": state}).json()
    assert inferred["resolvedMembers"][0]["kgm"] == 31.4
    assert inferred["netMembers"] > 0
    state["members"][0]["kgm"] = 0
    rejected = client.post("/api/calculate", json={"state": state}).json()
    assert rejected["resolvedMembers"][0]["kgm"] == 0
    assert rejected["netMembers"] == 0
    assert rejected["validation_errors"]


def test_csv_import_requires_verified_source_and_is_atomic(client):
    data = b"designation,family,kg/m,stock_lengths_m,supplier_rate_per_tonne\n100x100x5 SHS,SHS,14.9,8;12,2050\n"
    files = {"file": ("sections.csv", data, "text/csv")}
    assert client.post("/api/sections/import", files=files).status_code == 422
    response = client.post("/api/sections/import", files=files, data={"verified": "true", "source": "Supplier 2026 table"})
    assert response.status_code == 200, response.text
    assert response.json()["imported"] == 1
    assert response.json()["sections"][0]["stocks"] == [8, 12]
    invalid = data + b"BAD,SHS,-10,8,2050\n"
    assert client.post("/api/sections/import", files={"file": ("sections.csv", invalid)}, data={"verified": "true", "source": "Supplier"}).status_code == 422
    assert len(client.get("/api/settings").json()["sections"]) == 1


def test_excel_import_and_duplicate_designations():
    book = Workbook()
    ws = book.active
    ws.append(["section", "family", "kgm", "stocks", "verified", "source"])
    ws.append(["250UB31.4", "UB", 31.4, "9,12", True, "Designation mass / supplier"])
    stream = io.BytesIO()
    book.save(stream)
    result = import_sections(stream.getvalue(), "sections.xlsx")
    assert result[0]["kgm"] == 31.4
    assert result[0]["stocks"] == [9, 12]
    with pytest.raises(ValueError, match="duplicate"):
        import_sections(b"section,family,kgm,stocks,verified,source\n250UB31.4,UB,31.4,9,true,test\n250 UB 31.4,UB,31.4,9,true,test", "dupes.csv")


def test_reports_are_real_readable_artifacts_with_private_warnings(client, state):
    state["project"]["name"] = "=DANGEROUS()"
    state["project"]["boltEnabled"] = True
    state["members"][0]["mark"] = "=1+2"
    saved = client.post("/api/projects", json={"state": state}).json()
    assert saved["result"]["margin"] is None
    pid = saved["id"]
    response = client.get(f"/api/projects/{pid}/report.xlsx")
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    assert {"Summary", "Members", "Stock purchasing", "Stock layouts", "Plates and bolts", "Costs and hours", "Notes and drawings", "Rate snapshot"}.issubset(book.sheetnames)
    assert book["Members"]["A4"].value == "'=1+2"
    assert "Unverified" in str(book["Summary"]["B16"].value)
    response = client.get(f"/api/projects/{pid}/report.pdf")
    assert response.status_code == 200
    doc = fitz.open(stream=response.content, filetype="pdf")
    text = "\n".join(page.get_text() for page in doc)
    assert "INTERNAL ESTIMATE" in text
    assert "UNVERIFIED" in text
    assert "Confirm site dimensions" in text
    assert "250UB31.4" in text
    doc.close()


def test_reports_preserve_applied_plate_density_and_unresolved_drawing_evidence(client, state):
    state["project"]["quoteScope"] = "Selected canopy beams only"
    state["project"]["exclusions"] = "Exclude existing columns"
    state["rates"] = {"steel_density_kg_mm3": .000008}
    state["plateMode"] = "detailed"
    state["plates"] = [{"label": "BP1", "qty": 2, "length": 1000, "width": 500, "thick": 10}]
    state["pdfFlags"] = ["Drawing S101 revision D needs clarification"]
    state["pdfCrossReferences"] = {
        "review": "Scope and revision references need manual checking",
        "marks": [{"mark": "B1", "section_conflict": True, "references": [{"file_id": "drawing-1", "page": 1, "section": "250UB31.4", "evidence": "B1 beam schedule"}]}],
        "rl_suggestions": [{"mark": "C1", "suggested_length_mm": 3200, "review": "Verify column end conditions"}],
        "flags": ["Conflicting sections across revision"], "method": "Local schedule cross-reference",
    }
    saved = client.post("/api/projects", json={"state": state}).json()
    pid = saved["id"]
    book = load_workbook(io.BytesIO(client.get(f"/api/projects/{pid}/report.xlsx").content))
    assert book["Plates and bolts"]["G4"].value == pytest.approx(80)
    notes = "\n".join(str(cell.value) for row in book["Notes and drawings"] for cell in row)
    assert "S101 revision D" in notes
    assert "SECTION CONFLICT" in notes
    assert "PROVISIONAL RL height suggestion" in notes
    assert "3200" in notes
    assert "Selected canopy beams only" in notes
    assert "Exclude existing columns" in notes
    pdf = fitz.open(stream=client.get(f"/api/projects/{pid}/report.pdf").content, filetype="pdf")
    text = "\n".join(p.get_text() for p in pdf)
    assert "S101 revision D" in text
    assert "SECTION CONFLICT" in text
    assert "PROVISIONAL RL height suggestion" in text
    assert "Verify column end conditions" in text
    assert "Selected canopy beams only" in text
    assert "Exclude existing columns" in text
    pdf.close()


def test_backup_restore_keeps_database_drawings_and_previous_data(tmp_path, state):
    store = Store(tmp_path / "app_data")
    first = store.save(state, {"totalSell": 100})
    uploads = store.data_dir / "uploads"
    uploads.mkdir()
    (uploads / "drawing.pdf").write_bytes(b"original drawing")
    payload = store.backup()
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        assert {"manifest.json", Store.DB_NAME, "uploads/drawing.pdf"}.issubset(archive.namelist())
    store.save({**state, "project": {"name": "Later"}}, {"totalSell": 200}, first["id"], 1)
    (uploads / "drawing.pdf").write_bytes(b"later drawing")
    result = store.restore(payload)
    assert result["restored"] is True
    assert store.get(first["id"])["revision"] == 1
    assert (uploads / "drawing.pdf").read_bytes() == b"original drawing"
    previous = __import__("pathlib").Path(result["preserved_previous_data"])
    assert (previous / "uploads/drawing.pdf").read_bytes() == b"later drawing"


@pytest.mark.parametrize("unsafe", ["../escaped", "/absolute", "uploads/../../escaped", "uploads\\escaped", "C:/escaped"])
def test_restore_rejects_path_traversal_without_touching_data(tmp_path, state, unsafe):
    store = Store(tmp_path / "app_data")
    saved = store.save(state, {})
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr(unsafe, "hostile")
    with pytest.raises(ValueError):
        store.restore(stream.getvalue())
    assert store.get(saved["id"])["state"] == state
    assert not (tmp_path / "escaped").exists()


def test_restore_rejects_untrusted_database_and_bad_zip(tmp_path, state):
    store = Store(tmp_path / "app_data")
    saved = store.save(state, {})
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("manifest.json", json.dumps({"application": "SteelTakeoff AI", "format": 1}))
        archive.writestr(Store.DB_NAME, "not sqlite")
    with pytest.raises(ValueError):
        store.restore(stream.getvalue())
    with pytest.raises(ValueError):
        store.restore(b"not zip")
    assert store.get(saved["id"])["state"] == state


def test_invalid_excel_file_is_validation_error(client):
    response = client.post("/api/sections/import", files={"file": ("corrupt.xlsx", b"not a spreadsheet")}, data={"verified": "true", "source": "Supplier"})
    assert response.status_code == 422


def test_assembly_fields_and_fraction_alias_preserved(client):
    row = {"name": "Post", "section": "100x100x5 SHS", "labour_hours": 1.25, "efficiency": .9, "description": "Includes base plates"}
    result = client.put("/api/settings", json={"assemblies": [row]})
    assert result.status_code == 200
    saved = result.json()["assemblies"][0]
    assert saved["hours"] == 1.25
    assert saved["repeat_efficiency"] == 90
    assert saved["section"] == row["section"]
    assert saved["description"] == row["description"]
